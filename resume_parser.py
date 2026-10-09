"""
Resume Parser Module
Extracts text from PDF, DOCX, and TXT files.
"""
import re
from pathlib import Path
from typing import Optional

try:
    import fitz  # PyMuPDF
    HAS_PYMUPDF = True
except ImportError:
    HAS_PYMUPDF = False

try:
    from docx import Document
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


class ResumeParser:
    """Extract clean text from resume files."""

    def __init__(self):
        self.supported_exts = {".pdf", ".docx", ".txt", ".md"}

    def parse(self, file_path: str) -> str:
        """Parse a resume file and return clean text."""
        path = Path(file_path)
        ext = path.suffix.lower()

        if ext not in self.supported_exts:
            raise ValueError(f"Unsupported file type: {ext}. Use PDF, DOCX, or TXT.")

        if ext == ".pdf":
            return self._parse_pdf(path)
        elif ext == ".docx":
            return self._parse_docx(path)
        else:
            return self._parse_text(path)

    def _parse_pdf(self, path: Path) -> str:
        """Extract text from PDF using PyMuPDF."""
        if not HAS_PYMUPDF:
            raise ImportError("PyMuPDF not installed. Run: pip install PyMuPDF")

        text = ""
        with fitz.open(path) as doc:
            for page in doc:
                text += page.get_text()
        return self._clean_text(text)

    def _parse_docx(self, path: Path) -> str:
        """Extract text from DOCX."""
        if not HAS_DOCX:
            raise ImportError("python-docx not installed. Run: pip install python-docx")

        doc = Document(path)
        text = "\n".join([para.text for para in doc.paragraphs if para.text.strip()])
        return self._clean_text(text)

    def _parse_text(self, path: Path) -> str:
        """Read plain text file."""
        return self._clean_text(path.read_text(encoding="utf-8"))

    def _clean_text(self, text: str) -> str:
        """Clean extracted text."""
        # Remove excessive whitespace
        text = re.sub(r'\s+', ' ', text)
        # Remove special characters but keep meaningful ones
        text = re.sub(r'[^\w\s@.#+\-/|,:;()\[\]]', ' ', text)
        # Normalize bullets
        text = re.sub(r'[•●■▪◆►▸]', ' ', text)
        return text.strip()

    def parse_bytes(self, file_bytes: bytes, filename: str) -> str:
        """Parse from uploaded bytes (for Streamlit)."""
        ext = Path(filename).suffix.lower()

        if ext == ".pdf":
            if not HAS_PYMUPDF:
                raise ImportError("PyMuPDF not installed.")
            text = ""
            with fitz.open(stream=file_bytes, filetype="pdf") as doc:
                for page in doc:
                    text += page.get_text()
            return self._clean_text(text)
        elif ext == ".docx":
            if not HAS_DOCX:
                raise ImportError("python-docx not installed.")
            import io
            doc = Document(io.BytesIO(file_bytes))
            text = "\n".join([para.text for para in doc.paragraphs if para.text.strip()])
            return self._clean_text(text)
        elif ext in {".txt", ".md"}:
            return self._clean_text(file_bytes.decode("utf-8"))
        else:
            raise ValueError(f"Unsupported file type: {ext}")
